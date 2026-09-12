"""Unit test for Excel workbook generation."""

import openpyxl
import pytest
from src.network_analyzer import CurrencyConvertedBreakdown
from src.optimizer import (
    AssignmentRecord,
    CostBreakdown,
    FacilityStatus,
    InboundFlowRecord,
    NetworkKPIs,
    OptimizationResult,
)
from src.report_generator import generate_excel_report


def test_generate_excel_report():
    breakdown = CostBreakdown(1000.0, 500.0, 300.0, 400.0, 2200.0)
    converted = CurrencyConvertedBreakdown(
        currency="USD",
        symbol="$",
        exchange_rate=140.0,
        fixed_facility_cost=7.14,
        variable_handling_cost=3.57,
        inbound_freight_cost=2.14,
        outbound_freight_cost=2.86,
        total_landed_cost=15.71,
        cost_per_unit=0.0157,
    )
    result = OptimizationResult(
        solver_status="OPTIMAL",
        is_feasible=True,
        objective_value_dzd=2200.0,
        facilities=[
            FacilityStatus("F1", "DC Alger", True, 1000.0, 2000.0, 0.5, 1000.0, 500.0, 36.7, 3.0)
        ],
        assignments=[
            AssignmentRecord("R1", "Alger", "F1", "DC Alger", 1000.0, 1, 0.4, 400.0, 36.7, 3.0, 36.7, 3.0)
        ],
        inbound_flows=[
            InboundFlowRecord("S1", "Port", "F1", "DC Alger", 1000.0, 0.3, 300.0)
        ],
        cost_breakdown=breakdown,
        kpis=NetworkKPIs(1000.0, 1, 1.0, 2.2, 0.5),
        raw_status_code=1,
    )

    buf = generate_excel_report(result, converted)
    assert buf is not None
    assert buf.getbuffer().nbytes > 1000

    wb = openpyxl.load_workbook(buf)
    assert "Executive Summary" in wb.sheetnames
    assert "Distribution Centers" in wb.sheetnames
    assert "Customer Allocation" in wb.sheetnames
    assert "Inbound Supply Flows" in wb.sheetnames
    assert "Landed Cost Breakdown" in wb.sheetnames
